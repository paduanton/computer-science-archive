"""Gere o Word com código real e capturas, ou um rascunho identificado."""

import argparse
import json
import textwrap
from datetime import date, datetime

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from entrega_utils import ROOT, config, pending, read_json, report_inputs_hash


def paragraph_code(document, text):
    for line in text.splitlines():
        wrapped = textwrap.wrap(
            line.expandtabs(4), width=88, replace_whitespace=False,
            drop_whitespace=False, break_long_words=True, break_on_hyphens=False,
        ) or [""]
        for part in wrapped:
            document.add_paragraph(part, style="Codigo")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--rascunho", action="store_true")
    args = parser.parse_args()
    details = config()
    problems = pending()
    if problems and not args.rascunho:
        print("Relatório final ainda depende de:\n- " + "\n- ".join(problems))
        return 1
    identity = read_json("identificacao.json", {})
    validation = read_json("validacao.json", {})
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    for name in ("Title", "Heading 1", "Heading 2"):
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
        doc.styles[name].font.name = "Calibri"
        doc.styles[name].font.underline = False
    for border in list(doc.styles.element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
    code = doc.styles.add_style("Codigo", WD_STYLE_TYPE.PARAGRAPH)
    code.font.name, code.font.size = "Consolas", Pt(8.5)
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.line_spacing = 0.98
    code.paragraph_format.widow_control = False
    doc.core_properties.title = details["titulo"]
    doc.core_properties.author = identity.get("nome", "")
    doc.add_paragraph(details["titulo"], style="Title")
    doc.add_paragraph("Sistemas Distribuídos")
    if args.rascunho:
        doc.add_paragraph("RASCUNHO PARA REVISÃO", style="Heading 1")
        doc.add_paragraph("Este arquivo ainda não é a entrega final. Faltam as etapas listadas abaixo.")
        for problem in problems:
            doc.add_paragraph(problem, style="List Bullet")
    for key, label in (("nome", "Aluno"), ("matricula", "Matrícula"),
                       ("turma", "Turma"), ("professor", "Professor")):
        value = str(identity.get(key, "")).strip()
        if value:
            doc.add_paragraph(label + ": " + value)
    doc.add_paragraph("Data: " + date.today().strftime("%d/%m/%Y"))
    doc.add_heading("Objetivo", level=1)
    doc.add_paragraph(details["objetivo"])
    doc.add_heading("Ambiente e execução", level=1)
    doc.add_paragraph(
        "Execução em Linux usando uma imagem Python 3.12, com dependências fixadas. "
        "No computador Windows, Docker Desktop utiliza o backend WSL2. "
        "O README também descreve a execução direta em Linux com ambiente virtual. "
        "O servidor escuta nas interfaces do contêiner e a porta é publicada "
        "somente no endereço local do computador."
    )
    doc.add_paragraph("Dependências da aplicação e dos testes")
    paragraph_code(doc, (ROOT / "requirements.txt").read_text(encoding="utf-8"))
    doc.add_paragraph("Os comandos abaixo são executados na pasta do projeto.")
    if details["gerados"]:
        doc.add_paragraph(
            "A geração gRPC com Python requer o ambiente virtual com as dependências "
            "instaladas; o Docker também a executa durante o build."
        )
    for command in details["comandos"]:
        paragraph_code(doc, command)
    doc.add_heading("Implementação", level=1)
    for paragraph in details["implementacao"]:
        doc.add_paragraph(paragraph)
    doc.add_heading("Testes e resultados", level=1)
    doc.add_paragraph(details["testes"])
    if validation.get("aprovado"):
        validation_date = datetime.fromisoformat(validation["data_utc"]).strftime("%d/%m/%Y às %H:%M UTC")
        doc.add_paragraph(
            f"Na validação registrada em {validation_date}, "
            f"{validation['tests']} testes passaram em {validation['sistema']}, "
            f"com Python {validation['python']}, sem falhas, erros ou testes ignorados. "
            "O registro completo acompanha a entrega em validacao.txt."
        )
    doc.add_heading("Evidências de execução", level=1)
    if details.get("captura"):
        doc.add_paragraph(details["captura"])
    has_images = any((ROOT / "evidencias" / name).exists() for name, _ in details["imagens"])
    if has_images:
        landscape = doc.add_section(WD_SECTION.NEW_PAGE)
        landscape.orientation = WD_ORIENT.LANDSCAPE
        landscape.page_width, landscape.page_height = Cm(29.7), Cm(21)
    for index, (name, caption) in enumerate(details["imagens"], 1):
        image_path = ROOT / "evidencias" / name
        if image_path.exists():
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.page_break_before = index > 1
            paragraph.paragraph_format.keep_with_next = True
            paragraph.add_run(f"Figura {index}").bold = True
            figure = doc.add_picture(str(image_path), width=Cm(25))
            if figure.height > Cm(14.5):
                scale = Cm(14.5) / figure.height
                figure.width = int(figure.width * scale)
                figure.height = Cm(14.5)
            doc.add_paragraph(caption, style="Caption")
        else:
            doc.add_paragraph(f"Figura {index}: {caption}. Captura pendente: {name}.")
    if has_images:
        portrait = doc.add_section(WD_SECTION.NEW_PAGE)
        portrait.orientation = WD_ORIENT.PORTRAIT
        portrait.page_width, portrait.page_height = Cm(21), Cm(29.7)
    doc.add_heading("Referências", level=1)
    reference = ("https://grpc.io/docs/languages/python/" if details["gerados"]
                 else "https://websockets.readthedocs.io/")
    doc.add_paragraph(
        "Enunciado da atividade: enunciado.pdf. Documentação oficial: "
        + reference + " Procedimentos locais de execução: README.md."
    )
    doc.add_heading("Código fonte desenvolvido", level=1)
    doc.add_paragraph(
        "As listagens seguintes contêm a implementação do laboratório e seus testes. "
        "As quebras de linhas longas são apenas de apresentação. "
        "Os arquivos executáveis originais acompanham a entrega."
    )
    if details["gerados"]:
        doc.add_paragraph(
            "Os módulos calculator_pb2.py e calculator_pb2_grpc.py são produzidos "
            "automaticamente pelo compilador. Estão incluídos no projeto; "
            "o contrato autoral completo é apresentado a seguir."
        )
    for name in details["codigo"]:
        heading = doc.add_heading(name, level=2)
        heading.paragraph_format.keep_with_next = True
        paragraph_code(doc, (ROOT / name).read_text(encoding="utf-8"))
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    filename = "relatorio_rascunho.docx" if args.rascunho else "relatorio.docx"
    doc.save(ROOT / filename)
    if not args.rascunho:
        (ROOT / "relatorio_manifesto.json").write_text(
            json.dumps({"entradas_sha256": report_inputs_hash()}, indent=2) + "\n",
            encoding="utf-8",
        )
    print("Gerado:", ROOT / filename)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
